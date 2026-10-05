# -*- coding: utf-8 -*-
"""
Чтение присланных файлов с реквизитами: .txt, .pdf, .docx, .csv, .rtf.

Банковские реквизиты клиенты чаще шлют файлом, а не текстом. Из файла текст
достаётся точно, без всякого распознавания — это самый надёжный вход.

PDF бывает двух видов:
  * текстовый (выгрузка из банка) — текст лежит внутри, читается мгновенно;
  * скан или фото, сохранённое в PDF — текста внутри нет, нужен OCR страниц.
Второй случай обрабатывается, если установлен pymupdf или pdf2image.
"""

import io
import os
import re

SUPPORTED = (".txt", ".csv", ".log", ".md", ".rtf", ".pdf", ".docx")


def is_supported(filename, mime=""):
    ext = os.path.splitext(str(filename or ""))[1].lower()
    if ext in SUPPORTED:
        return True
    m = str(mime or "")
    return (m.startswith("text/") or "pdf" in m or "wordprocessing" in m)


def _decode(blob):
    """Подбираем кодировку: клиенты шлют и UTF-8, и windows-1251."""
    for enc in ("utf-8-sig", "utf-8", "cp1251", "cp866", "latin-1"):
        try:
            text = blob.decode(enc)
            if enc in ("utf-8-sig", "utf-8") or "�" not in text:
                return text
        except Exception:
            continue
    return blob.decode("utf-8", errors="replace")


def _strip_rtf(text):
    text = re.sub(r"\\'([0-9a-fA-F]{2})", lambda m: bytes([int(m.group(1), 16)]).decode("cp1251", "ignore"), text)
    text = re.sub(r"\\u(-?\d+)\s?\??", lambda m: chr(int(m.group(1)) % 65536), text)
    text = re.sub(r"\\par[d]?", "\n", text)
    text = re.sub(r"\\[a-zA-Z]+-?\d*\s?", "", text)
    return re.sub(r"[{}]", "", text)


def _read_pdf(blob):
    """Сначала пробуем вытащить текстовый слой, и только если его нет — OCR."""
    text = ""
    for mod, fn in (("pdfplumber", "plumber"), ("pypdf", "pypdf")):
        try:
            if fn == "plumber":
                import pdfplumber
                with pdfplumber.open(io.BytesIO(blob)) as pdf:
                    text = "\n".join((p.extract_text() or "") for p in pdf.pages[:10])
            else:
                from pypdf import PdfReader
                r = PdfReader(io.BytesIO(blob))
                text = "\n".join((p.extract_text() or "") for p in r.pages[:10])
            if text.strip():
                return text, "текстовый слой PDF", {}
        except Exception:
            continue

    pages = _pdf_to_images(blob)
    if not pages:
        return "", ("в PDF нет текстового слоя (это скан), а отрисовать страницы нечем — "
                    "поставьте pymupdf: pip install pymupdf"), {}
    try:
        import ocr_local
        res = ocr_local.read_images([(p, "image/png") for p in pages])
        src = f"скан PDF, распознано {res.get('parts', 0)} стр."
        if res.get("mrz"):
            src += ", прочитана машиночитаемая зона"
        return res.get("text", ""), src, res.get("mrz") or {}
    except Exception as e:
        return "", f"PDF — скан, распознать не удалось: {e}", {}


def _pdf_to_images(blob, max_pages=4, dpi=200):
    try:
        import fitz                                  # pymupdf
        doc = fitz.open(stream=blob, filetype="pdf")
        out = []
        for page in list(doc)[:max_pages]:
            pix = page.get_pixmap(dpi=dpi)
            out.append(pix.tobytes("png"))
        return out
    except Exception:
        pass
    try:
        from pdf2image import convert_from_bytes
        out = []
        for im in convert_from_bytes(blob, dpi=dpi)[:max_pages]:
            b = io.BytesIO(); im.save(b, "PNG"); out.append(b.getvalue())
        return out
    except Exception:
        return []


def _read_docx(blob):
    from docx import Document
    doc = Document(io.BytesIO(blob))
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            parts.append("\t".join(c.text for c in row.cells))
    return "\n".join(x for x in parts if x.strip())


def read_rich(blob, filename="", mime=""):
    """
    Разбирает присланный файл и сам определяет, что это.
    Возвращает {"text", "source", "mrz"}: mrz заполняется, если внутри оказался
    скан паспорта и удалось прочитать машиночитаемую зону.
    """
    text, source, mrz = _dispatch(blob, filename, mime)
    return {"text": text, "source": source, "mrz": mrz or {}}


def read(blob, filename="", mime=""):
    """Старый вызов: (текст, описание источника)."""
    text, source, _ = _dispatch(blob, filename, mime)
    return text, source


def _dispatch(blob, filename="", mime=""):
    ext = os.path.splitext(str(filename or ""))[1].lower()
    m = str(mime or "")

    if ext == ".pdf" or "pdf" in m:
        return _read_pdf(blob)

    if ext == ".docx" or "wordprocessing" in m:
        try:
            return _read_docx(blob), "документ Word", {}
        except Exception as e:
            return "", f"не читается .docx: {e}", {}

    if ext == ".rtf" or "rtf" in m:
        return _strip_rtf(_decode(blob)), "файл RTF", {}

    if ext in (".txt", ".csv", ".log", ".md") or m.startswith("text/"):
        return _decode(blob), "текстовый файл", {}

    return "", "такой тип файла я не читаю", {}
